# Coverage workbook (Aly asked: the coverage table as Excel, for us + the video description). Data = research/تحليل_التغطية.md.
import json, sys
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter

PROJ = Path(__file__).resolve().parents[2]
OUT = PROJ / 'output' / 'تغطية_كورس_بريمير.xlsx'
F = 'Arial'
OK, MISS, EXC = 'مغطي', 'ناقص', 'مستبعد'

# Adobe exam objectives (July 2025), verbatim titles from certifiedprofessional.adobe.com/premiere-pro (two long lines cut by the fetch tool, marked …)
EXAM = [
    ('1.1a', 'Determine whether content is relevant to the purpose, audience, and audience needs.', EXC, '', 'الجمهور (مش مونتاج)'),
    ('1.1b', 'Identify requirements based on how the project will be delivered, distributed, and/or consumed.', EXC, '', ''),
    ('1.1c', 'Demonstrate techniques for communicating ideas about project plans with peers and clients.', EXC, '', ''),
    ('1.1d', 'Demonstrate basic project management and workflow concepts.', EXC, '', ''),
    ('1.2a', 'Identify and implement captioning, subtitles, audio descriptions, and other assistive elements …', EXC, '', 'Accessibility'),
    ('1.2b', 'Apply techniques for clear and legible text overlays, proper color contrast, and audio clarity …', EXC, '', 'Accessibility'),
    ('1.3a', 'Determine legal and ethical considerations for using third-party content.', EXC, '', 'حقوق الملكية'),
    ('1.3b', 'Determine when and how to obtain permission to use images, audio, or video.', EXC, '', 'حقوق الملكية'),
    ('1.4a', 'Define video and audio terminology.', EXC, '1-5', 'اتشرح فعلا، بس الجزء كله مستبعد'),
    ('1.4b', 'Use visual storytelling techniques and editorial principles and guidelines.', EXC, '', ''),
    ('1.4c', 'Define common film form terms and principles.', EXC, '', ''),
    ('2.1a', 'Choose appropriate project settings.', OK, '6', ''),
    ('2.1b', 'Create a sequence that matches the delivery requirements.', OK, '8, 10', ''),
    ('2.2a', 'Identify, navigate, and manipulate elements of the Premiere interface.', OK, '6, 11', ''),
    ('2.2b', 'Organize and customize the workspace.', OK, '6, 46', ''),
    ('2.2c', 'Configure application preferences.', EXC, '', 'إعدادات البرنامج'),
    ('2.2d', 'Use the Button Editor.', EXC, '', 'إعدادات البرنامج'),
    ('2.3a', 'Controlling and working with the timeline and media.', OK, '11', ''),
    ('2.3b', 'Use markers.', OK, '16', ''),
    ('2.3c', 'Use guides, rulers, and safe margins.', MISS, '', ''),
    ('2.4a', 'Import media from various sources.', OK, '7', ''),
    ('3.1a', 'Organize assets in the Project panel.', OK, '7, 25', ''),
    ('3.2a', 'Configure sequences and clips in a project.', OK, '8, 10, 36', ''),
    ('3.2b', 'Configure tracks for playback and display', OK, '11, 23', ''),
    ('4.1a', 'Create a rough cut.', OK, '14', ''),
    ('4.2a', 'Create a transcript from a clip', MISS, '', 'Text-Based Editing (2023)'),
    ('4.2b', 'Modify videos using Text-Based Editing.', MISS, '', 'Text-Based Editing (2023)'),
    ('4.2c', 'Export transcripts.', MISS, '', 'Text-Based Editing (2023)'),
    ('4.3a', 'Add text and shapes to a sequence.', OK, '34', ''),
    ('4.3b', 'Use controls to modify text and graphic appearance.', OK, '34', ''),
    ('4.3c', 'Add and modify captions.', MISS, '', 'Captions'),
    ('4.4a', 'Resize clips.', OK, '26', ''),
    ('4.4b', 'Trim and refine clips.', OK, '18-21', ''),
    ('4.4c', 'Modify video clip settings.', OK, '20, 43', ''),
    ('4.5a', 'Use color correction methods and tools.', OK, '41, 42', ''),
    ('4.5b', 'Apply effect presets.', OK, '32', ''),
    ('4.5c', 'Apply effect presets on multiple clips.', OK, '40', 'Adjustment layer'),
    ('4.5d', 'Composite video clips.', OK, '28, 33, 45', 'Blend modes, Mask, Green screen'),
    ('4.6a', 'Modify attributes in the Effect Controls Panel and Properties Panel.', OK, '26', ''),
    ('4.6b', 'Apply and adjust video effects in the Timeline.', OK, '26-29, 38', ''),
    ('4.7a', 'Apply effects presets.', OK, '30', 'Audio crossfade بالـ Default transition'),
    ('4.7b', 'Add audio to a sequence.', OK, '37', ''),
    ('4.7c', 'Adjust audio on the timeline or using the Effect Controls panel and Essential Sound panel.', OK, '37, 38', ''),
    ('4.7d', 'Adjust audio.', OK, '37, 39', ''),
    ('4.7e', 'Record audio to the timeline.', OK, '12', ''),
    ('5.1a', 'Check a sequence for errors and project specifications.', MISS, '', ''),
    ('5.2a', 'Create a new version of an existing sequence using Auto Reframe.', MISS, '', 'Auto Reframe'),
    ('5.3a', 'Export a single frame, a clip, a range of a sequence, or an entire sequence.', OK, '24', ''),
    ('5.3b', 'Export using Adobe Media Encoder.', OK, '24', 'طريقة تانية للـ Export'),
    ('5.3c', 'Archive a project.', EXC, '', 'إعدادات البرنامج'),
]
DOMAIN = {'1': 'Working in the Video Industry', '2': 'Project Setup and Interface', '3': 'Organizing Video Projects', '4': 'Creating and Modifying Elements', '5': 'Publishing Digital Media'}

# Premiere user guide (helpx.adobe.com/premiere/desktop.html, Oct 2026): sub-sections -> status + course chapters
GUIDE = {
    'get-started': {'technical-requirements': EXC, 'download-and-install': EXC, 'preferences-and-settings': EXC, 'set-up-accessibility-features': EXC, 'keyboard-shortcuts': (OK, '14, 18'),
                    'set-up-directx-hdr-for-windows': EXC, 'use-touch-and-gesture-controls': EXC, 'tour-the-workspace': (OK, '6, 46'), 'customize-the-project-panel': (OK, '7'),
                    'source-and-program-monitor-adjustments': (OK, '11, 14')},
    'organize-media': {'create-projects': (OK, '6'), 'import-files': (OK, '7'), 'file-organization': (OK, '25'), 'apply-labeling': MISS, 'transfer-files': MISS, 'ingest-proxy-workflow': MISS, 'edit-metadata': MISS},
    'edit-projects': {'intro-to-editing': (OK, '13, 14'), 'edit-video-using-text-based-editing': MISS, 'trim-clips': (OK, '18-21'), 'change-clip-speed': (OK, '20, 43'), 'change-clip-sequence': (OK, '13-21'),
                      'modify-clip-properties': (OK, '10'), 'set-up-multi-camera-sequences-for-editing': MISS, 'compare-edits-to-source-clips': MISS, 'correct-mistakes': (OK, '10'), 'edit-vr-content': EXC,
                      'edit-with-generative-ai': MISS, 'edit-nested-sequences': (OK, '36')},
    'add-text-images': {'use-motion-graphics-templates': MISS, 'stylize-text': (OK, '34'), 'insert-images-and-graphics': (OK, '30, 35'), 'draw-objects': (OK, '34'), 'align-and-distribute-objects': (OK, '34'), 'insert-captions': MISS},
    'add-video-effects': {k: (OK, c) for k, c in [('types-of-effects', '32'), ('apply-video-effects', '32'), ('effects-and-transitions-library', '31, 32'), ('control-effects-and-transitions-using-keyframes', '26, 27, 29'),
                          ('create-masks-and-composites', '33'), ('commonly-used-effects', '32, 40, 41, 45'), ('apply-video-transitions', '30, 31'), ('work-with-masks', '33'), ('work-with-composites', '28, 45')]},
    'add-audio-effects': {'basic-audio-editing': (OK, '37'), 'advanced-audio-techniques': MISS, 'adjust-volume-and-levels': (OK, '37-39'), 'apply-audio-effects': MISS, 'apply-audio-transitions': (OK, '30'), 'use-adobe-stock-audio': EXC},
    'correct-color': {'color-correction-fundamentals': (OK, '42'), 'add-color-effects': (OK, '41, 42'), 'set-up-color-management': EXC, 'color-mode-fundamentals': MISS},
    'render-and-export': {'render-sequences-for-playback': (OK, '32'), 'export-files': (OK, '24'), 'stream-video': EXC},
    'collaborate-with-others': {k: EXC for k in ['collaborate-using-team-projects', 'share-for-review-using-frame-io', 'collaborate-using-creative-cloud-libraries', 'collaborate-using-productions']},
    'troubleshooting': {k: EXC for k in ['limitations-and-known-issues', 'media-issues', 'crash-issues', 'preferences-and-settings-issues', 'export-issues', 'audio-issues', 'playback-issues']},
    'use-premiere-with-other-apps': {'working-with-other-adobe-applications': MISS},
}
LEARN = [('Get started', '6, 7'), ('Sequences', '8, 10, 36'), ('Visual effects', '32, 45'), ('Color', '41, 42'), ('Audio', '37-39'), ('Keyframes', '4, 26, 27'), ('Graphics', '34, 35'), ('Titles', '34'), ('Animation', '26-29')]
SOURCES = [
    ('منهج امتحان Adobe (يوليو 2025)', 'https://certifiedprofessional.adobe.com/premiere-pro'),
    ('دليل Premiere الرسمي (نسخة 26.5 – سبتمبر 2026)', 'https://helpx.adobe.com/premiere/desktop.html'),
    ('مواضيع دروس Adobe الرسمية', 'https://www.adobe.com/learn/premiere-pro'),
    ('Text-Based Editing اتعلن أبريل 2023', 'https://blog.adobe.com/en/publish/2023/04/13/nab-2023-introducing-text-based-editing-premiere-pro-properties-panel-after-effects-more'),
    ('الكورس الأصلي (2018)', 'https://www.youtube.com/playlist?list=PLiIobAzklBzFqp6mlEFmdYyUxcIXkXJ3y'),
]

tree = json.load(open(PROJ / 'research/user_guide_tree.json', encoding='utf-8'))
thin = Side(style='thin', color='D0D5DD')
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
HEAD = PatternFill('solid', fgColor='1F4E9F')
FILL = {OK: PatternFill('solid', fgColor='D8F5E3'), MISS: PatternFill('solid', fgColor='FFE1DD'), EXC: PatternFill('solid', fgColor='ECEEF2')}


def sheet(wb, title, headers, rows, widths):
    ws = wb.create_sheet(title)
    ws.sheet_view.rightToLeft = True
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(name=F, bold=True, color='FFFFFF'); c.fill = HEAD; c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True); c.border = BORDER
    for r in rows:
        ws.append(list(r))
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = Font(name=F); c.border = BORDER; c.alignment = Alignment(vertical='center', wrap_text=True)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = 'A2'
    return ws


def color_status(ws, col, n):
    rng = f'{col}2:{col}{n + 1}'
    for k, fill in FILL.items():
        ws.conditional_formatting.add(rng, CellIsRule(operator='equal', formula=[f'"{k}"'], fill=fill))


wb = Workbook()
wb.remove(wb.active)
summary = wb.create_sheet('الملخص')

ex = sheet(wb, 'منهج الامتحان', ['الكود', 'الجزء', 'النقطة (نص Adobe)', 'الحالة', 'فصول الكورس', 'ملاحظة'],
           [(c, f'{c[0]} · {DOMAIN[c[0]]}', t, s, ch, n) for c, t, s, ch, n in EXAM], [8, 30, 70, 11, 14, 34])
color_status(ex, 'D', len(EXAM))

grows = []
for sec, subs in tree.items():
    for sub, arts in subs.items():
        v = GUIDE.get(sec, {}).get(sub, EXC)
        st, ch = (v if isinstance(v, tuple) else (v, ''))
        grows.append((sec, sub, len(arts), st, ch))
gd = sheet(wb, 'دليل Premiere 2026', ['القسم', 'الجزء', 'عدد المقالات', 'الحالة', 'فصول الكورس'], grows, [26, 46, 12, 11, 16])
color_status(gd, 'D', len(grows))

ln = sheet(wb, 'مواضيع Adobe Learn', ['الموضوع', 'الحالة', 'فصول الكورس'], [(t, OK, ch) for t, ch in LEARN], [24, 11, 18])
color_status(ln, 'B', len(LEARN))

src = sheet(wb, 'المصادر', ['المصدر', 'اللينك'], SOURCES, [44, 90])

# summary with live formulas
s = summary
s.sheet_view.rightToLeft = True
s['A1'] = 'الكورس ده مغطي اد إيه؟ — كورس Premiere (علي شرش، 2018) مقارنة ببريمير 2026'
s['A1'].font = Font(name=F, bold=True, size=14)
s['A2'] = 'القاعدة: لو الأداة اتشرحت تتحسب كلها، حتى لو مش كل خصائصها اتشرحت. المستبعد مش مونتاج (الجمهور، حقوق الملكية، Accessibility، الإعدادات، التسطيب، الشغل الجماعي، حل المشاكل).'
s['A2'].font = Font(name=F, italic=True, color='555555')
s.append([])
s.append(['المقياس', 'مغطي', 'ناقص', 'مستبعد', 'النسبة (مغطي ÷ (مغطي + ناقص))'])
for c in s[4]:
    c.font = Font(name=F, bold=True, color='FFFFFF'); c.fill = HEAD; c.alignment = Alignment(horizontal='center', wrap_text=True); c.border = BORDER
for i, (name, sh, col) in enumerate([('مواضيع دروس Adobe', 'مواضيع Adobe Learn', 'B'), ('النقاط التقنية في منهج الامتحان', 'منهج الامتحان', 'D'), ('أقسام المونتاج في دليل Premiere 2026', 'دليل Premiere 2026', 'D')]):
    r = 5 + i
    s[f'A{r}'] = name
    for j, (k, L) in enumerate([(OK, 'B'), (MISS, 'C'), (EXC, 'D')]):
        s[f'{L}{r}'] = f"=COUNTIF('{sh}'!{col}:{col},\"{k}\")"
    s[f'E{r}'] = f'=IFERROR(B{r}/(B{r}+C{r}),0)'
    s[f'E{r}'].number_format = '0%'
    for c in s[r]:
        c.font = Font(name=F, bold=c.column == 5); c.border = BORDER; c.alignment = Alignment(horizontal='center' if c.column > 1 else 'right')
s['A9'] = 'جزء المجال (الجمهور، الحقوق، Accessibility) من منهج الامتحان:'
s['E9'] = "=COUNTIF('منهج الامتحان'!B:B,\"1 · Working in the Video Industry\")/COUNTA('منهج الامتحان'!A2:A51)"
s['E9'].number_format = '0%'
s['A10'] = 'المصادر في شيت "المصادر". تحليل Claude — 17 أكتوبر 2026.'
for r in (9, 10):
    s[f'A{r}'].font = Font(name=F, color='555555')
s['E9'].font = Font(name=F, bold=True)
for col, w in zip('ABCDE', [42, 10, 10, 10, 22]):
    s.column_dimensions[col].width = w
OUT.parent.mkdir(exist_ok=True)
wb.save(OUT)
print(OUT)
