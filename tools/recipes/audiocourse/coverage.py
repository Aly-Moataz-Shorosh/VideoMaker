# Claude's coverage of the audio course (Aly 2026-10-25: "رأي كلود" = how many % of each yardstick the course covers).
# Yardsticks = TODAY's official references (checked 2026-10-25). Fundamentals + mics at their own sections, programs one row per TOOL
# (Aly: 89% for Audition by chapters felt too high) (Aly: if he explained
# the bigger topic, its sub-details are not held against him; true numbers only, and write on screen what was left out).
# Each section: (name, covered, where in the session / what is inside). ex = what was taken out of the count, shown on screen.
#   python coverage.py   -> work/coverage.xlsx (for the description)
COVER = [
    dict(head='أساسيات الصوت الديجيتال', en='Digital audio fundamentals (Adobe: Understanding sound + Digitizing audio)', src='Adobe Audition User Guide',
         urls=['https://helpx.adobe.com/audition/desktop/digital-audio-fundamentals/sound.html',
               'https://helpx.adobe.com/audition/desktop/digital-audio-fundamentals/digitizing-audio.html'],
         miss='تداخل الموجات، حجم الملف', ex='dBFS لأنه جزء من Bit depth، وطريقة Audition جوه البرنامج', items=[
             ('What sound is (air pressure waves)', 1, '22:31 molecules push each other'),
             ('Waveform measurements (amplitude, cycle, frequency, wavelength; phase not said)', 1, '24:37 - 30:06, 36:10, 109:26'),
             ('How sound waves interact', 0, ''),
             ('Digitizing audio (analog -> numbers)', 1, '26:06 quantization'),
             ('Sample rate', 1, '26:34 - 31:52, 100:31, 114:42 48 kHz'),
             ('Bit depth', 1, '28:32 - 37:33, 100:40'),
             ('Audio file contents and size', 0, ''),
         ]),
    dict(head='المايكات', en='Microphones (Shure polar patterns + Sweetwater mic types)', src='Shure + Sweetwater',
         urls=['https://www.shure.com/en-US/insights/microphone-directionality-polar-pattern-basics',
               'https://www.sweetwater.com/insync/dynamic-vs-condenser-vs-ribbon-microphones-how-are-they-different/'],
         has='وزيادة: Shotgun، Pop filter، Dead cat', miss='Figure-8', ex='Super/Hyper cardioid لأن Shure بيحطهم تحت Unidirectional', items=[
             ('Dynamic', 1, '39:28'), ('Condenser', 1, '40:12'), ('Ribbon', 1, '39:34'),
             ('Omnidirectional', 1, '42:52 / 44:33'),
             ('Unidirectional (cardioid; super / hyper cardioid are its sub-types)', 1, '42:10 / 43:09 cardioid, 43:20 shotgun'),
             ('Bi-directional (figure-8)', 0, ''),
         ]),
    dict(head='الصوت في Premiere Pro', en='Audio in Premiere Pro (User Guide chapter "Add audio effects", one row per tool)', src='Adobe Premiere User Guide',
         urls=['https://helpx.adobe.com/premiere/desktop.html  (chapter: Add audio effects)'], unit='أدوات',
         has='Essential Sound بكل اللي جواه',
         # Aly review 3: the full Premiere course (youtu.be/D4fqs2c0DfY) = thumbnail + QR; thumb = HIS thumbnail of that video
         link=dict(text='كورس Premiere الكامل', thumb='projects/audio/pics/premiere_course_thumb.jpg', qr='projects/audio/qr/premiere_course.png'), miss='Transitions، Mixer، Loudness، Voice-over', ex='Adobe Stock، وحاجات صغيرة زي Bulk mute', items=[
             ('Camera audio vs mic audio on separate tracks', 1, '73:24'),
             ('Synchronize audio and video', 1, '75:22 clap sync'),
             ('Volume, gain and keyframes', 1, '77:40, 92:27'),
             ('Mute / Solo', 1, '73:50, 78:34'),
             ('Essential Sound panel (types + presets; Repair, Enhance Speech, Ducking, Remix, auto tag live inside it)', 1, '80:24 - 82:44, 86:29'),
             ('Audio effects (library, apply, adjust)', 1, '79:18, 89:08'),
             ('Audio transitions (fades, crossfades)', 0, 'named only at 79:30'),
             ('Audio Track Mixer, submix, pan / balance, automation', 0, ''),
             ('Audio channels (stereo / mono mapping)', 0, 'stereo / mono explained in Audition only'),
             ('Loudness meter, monitor levels, auto-match loudness', 0, ''),
             ('Record audio / voice-over', 0, 'recorded in Audition instead'),
         ]),
    dict(head='Adobe Audition', en='Adobe Audition (User Guide, one row per tool)', src='Adobe Audition User Guide',
         urls=['https://helpx.adobe.com/audition/desktop/digital-audio-fundamentals/digitizing-audio.html  (side menu = whole guide)'], unit='أدوات',
         miss='Multitrack، Essential Sound، Loudness، Markers', ex='المقدمة، الاختصارات، Surround 5.1، Undo', items=[
             ('Workspace and panels (Window menu)', 1, '96:57 - 97:46'),
             ('Import files (Files panel)', 1, '99:16'),
             ('Navigate, zoom, play', 1, '100:11 - 100:27'),
             ('Record audio (new file: sample rate, stereo / mono, bit depth)', 1, '113:56 - 117:11'),
             ('Monitor recording and playback levels', 0, ''),
             ('Waveform Editor + select audio (samples up close)', 1, '100:27 - 101:01, 101:28, 104:33'),
             ('Copy, cut, paste, delete audio', 0, ''),
             ('Visual fades and amplitude', 0, ''),
             ('Markers', 0, ''),
             ('Spectral display / analyze phase, frequency, amplitude', 0, ''),
             ('Essential Sound panel', 0, ''),
             ('Match loudness', 0, ''),
             ('Reduce noise and restore audio (noise print)', 1, '102:41 - 108:13'),
             ('Apply effects + presets (reverb, delay / echo, filter, amplitude families)', 1, '109:26 - 113:18'),
             ('Multitrack Editor (mixing, envelopes, routing, remix)', 0, ''),
             ('Video in Audition (import video, its audio)', 1, '99:41'),
             ('Save and export', 1, '117:21 - 118:04'),
         ]),
]
for c in COVER:
    c['n'] = sum(x[1] for x in c['items']); c['of'] = len(c['items']); c['v'] = round(c['n'] / c['of'], 4)


def cards():
    return [dict(head=c['head'], v=c['v'], n=c['n'], of=c['of'], unit=c.get('unit', 'أقسام'), has=c.get('has'), miss=c['miss'], ex=c['ex'], src=c['src'], link=c.get('link')) for c in COVER]


if __name__ == '__main__':
    from pathlib import Path
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook(); ws = wb.active; ws.title = 'Coverage'
    ws.append(['Yardstick', 'Section', 'Covered', 'Where in the session (min:sec)', 'Source / left out of the count'])
    for c in ws[1]: c.font = Font(bold=True)
    for c in COVER:
        ws.append([f"{c['en']}: {c['n']}/{c['of']} = {round(c['v'] * 100)}%", '', '', '', ' | '.join(c['urls'])])
        ws.cell(ws.max_row, 1).font = Font(bold=True)
        for name, ok, where in c['items']:
            ws.append(['', name, 'yes' if ok else 'no', where, ''])
            ws.cell(ws.max_row, 3).fill = PatternFill('solid', fgColor='C8F0D0' if ok else 'F8D0C8')
        ws.append(['', 'Left out of the count', '', '', c['ex']])
    for col, w in zip('ABCDE', (40, 70, 10, 44, 90)): ws.column_dimensions[col].width = w
    out = Path(__file__).resolve().parent.parent / 'coverage.xlsx'
    wb.save(out)
    for c in COVER: print(c['en'][:30], c['n'], '/', c['of'], round(c['v'] * 100), '%')
    print(out)
